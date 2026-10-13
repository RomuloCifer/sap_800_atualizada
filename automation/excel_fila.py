"""
Fila do Excel: primeira linha com Status = Pendente.

Colunas por nome do cabeçalho (não pela letra):
  Status | Batch | Issuer Id | Total Value
"""

from __future__ import annotations

import os
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional, Union

ROOT = Path(__file__).resolve().parent.parent
EXCEL_DIR = ROOT / "excel"

STATUS_ALVO = "pendente"
STATUS_OK = "A-enviado"
STATUS_ERRO = "A-erro"

# Nomes aceitos no cabeçalho (case-insensitive)
COL_STATUS = ("status",)
COL_BATCH = ("batch",)
COL_ISSUER = ("issuer id", "issuer_id", "issuerid")
COL_TOTAL = ("total value", "total_value", "totalvalue")
COL_CBS_IBS = ("cbs/ibs", "cbs ibs", "cbs_ibs", "cbsibs")


@dataclass
class LinhaFila:
    caminho: Path
    linha: int  # 1-based na planilha
    status: str
    batch: str
    issuer_id: str
    total_value: Optional[str] = None
    cbs_ibs: Optional[str] = None

    def to_contexto(self):
        # type: () -> dict
        ctx = {
            "batch": self.batch,
            "issuer_id": self.issuer_id,
            "excel_linha": self.linha,
            "excel_arquivo": str(self.caminho),
            "excel_status": self.status,
            "cbs_ibs": self.cbs_ibs or "",
        }
        if self.total_value is not None:
            ctx["excel_total_value"] = self.total_value
        return ctx


def _norm_header(valor):
    # type: (object) -> str
    if valor is None:
        return ""
    return " ".join(str(valor).strip().lower().split())


def _cell_str(valor):
    # type: (object) -> str
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return str(valor)
    if isinstance(valor, float):
        if valor == int(valor):
            return str(int(valor))
        return str(valor).strip()
    if isinstance(valor, int):
        return str(valor)
    return str(valor).strip()


def _achar_coluna(headers, candidatos):
    # type: (List[str], tuple) -> Optional[int]
    for i, h in enumerate(headers):
        if h in candidatos:
            return i
    return None


def resolver_arquivo_excel(caminho=None):
    # type: (Union[str, Path, None]) -> Path
    if caminho:
        p = Path(caminho)
        if not p.is_file():
            raise FileNotFoundError("Excel nao encontrado: {}".format(p))
        return p

    if not EXCEL_DIR.is_dir():
        raise FileNotFoundError("Pasta excel/ nao existe: {}".format(EXCEL_DIR))

    candidatos = sorted(
        list(EXCEL_DIR.glob("*.xlsx")) + list(EXCEL_DIR.glob("*.XLSX")),
        key=lambda x: x.stat().st_mtime,
        reverse=True,
    )
    # Evita duplicata case-insensitive no Windows
    vistos = set()
    unicos = []
    for c in candidatos:
        key = str(c.resolve()).lower()
        if key in vistos:
            continue
        vistos.add(key)
        unicos.append(c)

    if not unicos:
        raise FileNotFoundError(
            "Nenhum .xlsx em {}. Coloque a planilha la.".format(EXCEL_DIR)
        )
    return unicos[0]


def _load_workbook(path, **kwargs):
    # type: (Path, object) -> object
    """Abre o xlsx; se estiver corrompido, mensagem clara (não o KeyError do zip)."""
    import openpyxl

    try:
        return openpyxl.load_workbook(path, **kwargs)
    except KeyError as exc:
        if "Content_Types" in str(exc):
            raise RuntimeError(
                "Excel corrompido ou incompleto: '{}'.\n"
                "Substitua por uma cópia boa em excel/ e feche o arquivo no Excel "
                "antes de gravar Status.".format(Path(path).name)
            ) from exc
        raise
    except zipfile.BadZipFile as exc:
        raise RuntimeError(
            "Excel inválido (não é .xlsx ZIP): '{}'.\n"
            "Salve de novo como .xlsx e coloque em excel/.".format(Path(path).name)
        ) from exc


def _salvar_workbook_atomico(wb, path):
    # type: (object, Path) -> None
    """
    Grava em arquivo temporário na mesma pasta e só depois troca o original
    (os.replace). Assim um Ctrl+C / crash no meio do save não corrompe o xlsx.
    """
    path = Path(path)
    fd, tmp_name = tempfile.mkstemp(
        prefix="._{}.".format(path.stem[:40]),
        suffix=".xlsx.tmp",
        dir=str(path.parent),
    )
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        wb.save(str(tmp_path))
        # Confirma que o ZIP ficou completo antes de substituir
        with zipfile.ZipFile(str(tmp_path), "r") as zf:
            nomes = zf.namelist()
        if "[Content_Types].xml" not in nomes:
            raise RuntimeError(
                "Falha ao gravar Excel: arquivo temporário incompleto "
                "(sem [Content_Types].xml). Original NÃO foi alterado."
            )
        os.replace(str(tmp_path), str(path))
        tmp_path = None  # já foi renomeado; não apagar
    except BaseException:
        # KeyboardInterrupt inclusive: original intacto
        if tmp_path is not None:
            try:
                if tmp_path.is_file():
                    tmp_path.unlink()
            except Exception:
                pass
        raise


def primeira_pendente(caminho=None):
    # type: (Union[str, Path, None]) -> LinhaFila
    """
    Retorna a primeira linha cujo Status (normalizado) == 'pendente'.
    Nao pega 'Enviados ... pendente revisao' — so Status exatamente Pendente.
    """
    path = resolver_arquivo_excel(caminho)
    wb = _load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            raise RuntimeError("Planilha vazia: {}".format(path))

        headers = [_norm_header(h) for h in header_row]
        i_status = _achar_coluna(headers, COL_STATUS)
        i_batch = _achar_coluna(headers, COL_BATCH)
        i_issuer = _achar_coluna(headers, COL_ISSUER)
        i_total = _achar_coluna(headers, COL_TOTAL)
        i_cbs = _achar_coluna(headers, COL_CBS_IBS)

        faltando = []
        if i_status is None:
            faltando.append("Status")
        if i_batch is None:
            faltando.append("Batch")
        if i_issuer is None:
            faltando.append("Issuer Id")
        if faltando:
            raise RuntimeError(
                "Cabecalhos nao encontrados em {}: {}. "
                "Cabecalhos lidos: {}".format(
                    path.name, ", ".join(faltando), ", ".join(headers[:25])
                )
            )

        for num, row in enumerate(rows, start=2):
            if row is None:
                continue
            status = _cell_str(row[i_status] if i_status < len(row) else None)
            if status.lower() != STATUS_ALVO:
                continue
            batch = _cell_str(row[i_batch] if i_batch < len(row) else None)
            issuer = _cell_str(row[i_issuer] if i_issuer < len(row) else None)
            if not batch:
                raise RuntimeError(
                    "Linha {} Status=Pendente sem Batch em {}.".format(num, path.name)
                )
            if not issuer:
                raise RuntimeError(
                    "Linha {} Status=Pendente sem Issuer Id em {}.".format(
                        num, path.name
                    )
                )
            total = None
            if i_total is not None and i_total < len(row):
                total = _cell_str(row[i_total]) or None
            cbs = None
            if i_cbs is not None and i_cbs < len(row):
                cbs = _cell_str(row[i_cbs]) or None
            return LinhaFila(
                caminho=path,
                linha=num,
                status=status,
                batch=batch,
                issuer_id=issuer,
                total_value=total,
                cbs_ibs=cbs,
            )
    finally:
        wb.close()

    raise RuntimeError(
        "Nenhuma linha com Status='Pendente' em {}.".format(path.name)
    )


def ler_cbs_ibs_linha(caminho, linha_1based):
    # type: (Union[str, Path], int) -> str
    """Relê CBS/IBS de uma linha específica (1-based, com cabeçalho na 1)."""
    path = Path(caminho)
    wb = _load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header_row = next(rows)
        headers = [_norm_header(h) for h in header_row]
        i_cbs = _achar_coluna(headers, COL_CBS_IBS)
        if i_cbs is None:
            return ""
        alvo = int(linha_1based)
        for num, row in enumerate(rows, start=2):
            if num != alvo:
                continue
            if row is None or i_cbs >= len(row):
                return ""
            return _cell_str(row[i_cbs])
    finally:
        wb.close()
    return ""


def atualizar_status(caminho, linha_1based, novo_status):
    # type: (Union[str, Path], int, str) -> None
    """
    Grava Status na linha (1-based).
    Save atômico (temp → replace): Ctrl+C no meio não corrompe o original.
    Feche o Excel na planilha antes — arquivo aberto trava a gravação.
    """
    path = Path(caminho)
    if not path.is_file():
        raise FileNotFoundError("Excel nao encontrado: {}".format(path))

    wb = _load_workbook(path)
    try:
        ws = wb.active
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [_norm_header(h) for h in header_row]
        i_status = _achar_coluna(headers, COL_STATUS)
        if i_status is None:
            raise RuntimeError("Coluna Status nao encontrada em {}.".format(path.name))
        ws.cell(row=int(linha_1based), column=i_status + 1, value=str(novo_status))
        _salvar_workbook_atomico(wb, path)
    finally:
        wb.close()


def tem_pendente(caminho=None):
    # type: (Union[str, Path, None]) -> bool
    try:
        primeira_pendente(caminho)
        return True
    except RuntimeError:
        return False


def marcar_resultado_do_contexto(contexto, sucesso, status_custom=None):
    # type: (dict, bool, Optional[str]) -> Optional[str]
    """
    Se o contexto veio do Excel (excel_arquivo + excel_linha), grava Status.
    sucesso True → A-enviado; False → A-erro.
    Se status_custom for informado (ex.: motivo do Parar), grava esse texto.
    Retorna o status gravado ou None se não havia linha Excel.
    """
    ctx = contexto or {}
    caminho = ctx.get("excel_arquivo")
    linha = ctx.get("excel_linha")
    if not caminho or not linha:
        return None
    if status_custom is not None and str(status_custom).strip():
        status = str(status_custom).strip()
    else:
        status = STATUS_OK if sucesso else STATUS_ERRO
    atualizar_status(caminho, int(linha), status)
    return status
